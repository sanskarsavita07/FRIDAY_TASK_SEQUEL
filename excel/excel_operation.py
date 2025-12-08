# excel_operation.py

import openpyxl
from config.config import (
    FILE_PATH
)
import warnings

def clearWarning():
    """Suppress openpyxl warnings"""
    warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

def loadExcel(filePath):
    """Load Excel workbook and return workbook and active sheet objects"""
    try:
        workBook = openpyxl.load_workbook(filePath, data_only=True)
        sheet = workBook.active
        print("Excel file loaded successfully")
        return workBook, sheet
    except FileNotFoundError:
        print(f"Error: File not found at {filePath}")
        raise
    except Exception as e:
        print(f"Error loading Excel file: {e}")
        raise

def extractMainData(sheet):
    """Extract main table data from Excel sheet"""
    try:
        mainData = []
        for row in sheet.iter_rows(
                min_row=6,
                max_row=101,
                min_col=1,
                max_col=8,
                values_only=True):
            if any(row):
                mainData.append(row)
                print(row)
        print(f"Extracted {len(mainData)} rows (Main Table)")
        print(mainData)
        return mainData
    except Exception as e:
        print(f"Error extracting main data: {e}")
        raise

def extractSummaryData(sheet):
    """Extract summary table data from Excel sheet"""
    try:
        summaryData = []
        for row in sheet.iter_rows(
                min_row=103,
                max_row=106,
                min_col=1,
                max_col=7,
                values_only=True):
            if any(row):
                summaryData.append(row)
                print(row)
        print(f"Extracted {len(summaryData)} rows (Summary Table)")
        print(summaryData)
        return summaryData
    except Exception as e:
        print(f"Error extracting summary data: {e}")
        raise

def closeWorkbook(workBook):
    """Close the Excel workbook"""
    try:
        if workBook:
            workBook.close()
            print("Workbook closed successfully")
    except Exception as e:
        print(f"Error closing workbook: {e}")